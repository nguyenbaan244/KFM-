"""
Script chuyển đổi dữ liệu từ File PO sang File Nhập Hàng theo định dạng chuẩn mới.
Quy tắc ánh xạ (Mapping):
  - orderInboundCode (Cột A)    <- Cột B PO (Mã PO)
  - productCode (Cột B)         <- Cột T PO (Mã hàng)
  - expectedQuantity (Cột C)    <- Cột AD PO (Số lượng PR thực nhận) [Nếu = 0 thì bỏ qua]
  - estimateReceiveTime (Cột D) <- Cột G PO (Ngày giao hàng NCC xác nhận, định dạng Text dd/mm/yyyy - nếu nhỏ hơn today thì đổi thành today)
  - customerNote (Cột E)        <- Để trống
  - zoneType (Cột F)            <- Điền B2B
  - supplier (Cột G)            <- Cột K PO (Tên nhà cung cấp)
  - productionDate (Cột H)      <- Cột X PO (NSX)
  - expiryDate (Cột I)          <- Cột Y PO (HSD)
  - inboundDate (Cột J)         <- Today (Ngày hiện tại dạng dd/mm/yyyy)

Sử dụng:
    python convert_po.py [duong_dan_file_po] [duong_dan_file_xuat] [duong_dan_file_template]
Ví dụ:
    python convert_po.py PO.xlsx ket_qua_nhap_hang.xlsx
"""

import sys
import os
import datetime
import openpyxl

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

def parse_date(val):
    if val is None:
        return None
    if isinstance(val, datetime.datetime):
        return val.date()
    if isinstance(val, datetime.date):
        return val
    if isinstance(val, (int, float)):
        if val > 30000:
            try:
                import openpyxl.utils.datetime
                return openpyxl.utils.datetime.from_excel(val).date()
            except Exception:
                pass
    s = str(val).strip()
    if not s or s.lower() == "none":
        return None
    if " " in s:
        s = s.split()[0]
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%Y/%m/%d", "%d/%m/%y", "%d-%m-%y"):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None

def format_date_str(val):
    if val is None:
        return ""
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime("%d/%m/%Y")
    s = str(val).strip()
    if s.lower() == "none":
        return ""
    if " " in s:
        parts = s.split()
        if "/" in parts[0] or "-" in parts[0]:
            return parts[0]
    return s

def process_estimate_date(val, today=None):
    if today is None:
        today = datetime.date.today()
    if isinstance(today, datetime.datetime):
        today = today.date()
    if val is None:
        return ""
    parsed = parse_date(val)
    if parsed is not None:
        if parsed < today:
            return today.strftime("%d/%m/%Y")
        return parsed.strftime("%d/%m/%Y")
    return format_date_str(val)

def find_column_indices(ws_po):
    # Default indices based on user spec:
    # B: 2 (Mã PO)
    # T: 20 (Mã hàng)
    # AD: 30 (Số lượng PR thực nhận)
    # G: 7 (Ngày giao hàng NCC xác nhận)
    # K: 11 (Tên NCC)
    # X: 24 (NSX)
    # Y: 25 (HSD)
    cols = {
        'po': 2,
        'product': 20,
        'qty': 30,
        'est_date': 7,
        'supplier': 11,
        'nsx': 24,
        'hsd': 25
    }

    # Header scan (row 1)
    for col in range(1, ws_po.max_column + 1):
        raw = str(ws_po.cell(1, col).value or '').strip().lower()
        if raw in ['mã po', 'ma po']:
            cols['po'] = col
        elif raw in ['mã hàng', 'ma hang', 'mã sp', 'ma sp', 'barcode', 'mã vạch', 'ma vach']:
            cols['product'] = col
        elif 'số lượng pr (thực nhận)' in raw or 'sl thực nhận' in raw or 'so luong thuc nhan' in raw or 'số lượng thực nhận' in raw:
            cols['qty'] = col
        elif 'ngày giao hàng ncc xác nhận' in raw or 'ngay giao hang ncc xac nhan' in raw or 'xác nhận giao hàng' in raw:
            cols['est_date'] = col
        elif raw in ['tên nhà cung cấp', 'ten nha cung cap', 'tên ncc', 'ten ncc']:
            cols['supplier'] = col
        elif raw in ['nsx', 'ngày sản xuất', 'ngay san xuat']:
            cols['nsx'] = col
        elif raw in ['hsd', 'hạn sử dụng', 'han su dung']:
            cols['hsd'] = col

    return cols

def convert_po_to_inbound(po_file_path=None, template_path=None, output_path=None, zone="B2B", note=""):
    # Determine default paths
    if not po_file_path:
        for candidate in ["PO.xlsx", "Export-PO-Data.xlsx"]:
            if os.path.exists(candidate):
                po_file_path = candidate
                break
        if not po_file_path:
            po_file_path = "PO.xlsx"

    if not template_path:
        for candidate in ["file_nhap_hang (1) (1).xlsx", "file_nhap_hang.xlsx"]:
            if os.path.exists(candidate):
                template_path = candidate
                break
        if not template_path:
            template_path = "file_nhap_hang (1) (1).xlsx"

    if not os.path.exists(po_file_path):
        print(f"Lỗi: Không tìm thấy file PO '{po_file_path}'")
        return False
    
    if not os.path.exists(template_path):
        print(f"Lỗi: Không tìm thấy file template '{template_path}'")
        return False

    print(f"Đang đọc file PO: {po_file_path}...")
    wb_po = openpyxl.load_workbook(po_file_path, data_only=True)
    ws_po = wb_po.active

    col_map = find_column_indices(ws_po)

    print(f"Đang đọc file mẫu: {template_path}...")
    wb_out = openpyxl.load_workbook(template_path)
    if 'File nhập hàng' in wb_out.sheetnames:
        ws_out = wb_out['File nhập hàng']
    else:
        ws_out = wb_out.active

    # Xóa các dòng mẫu từ dòng 3 trở đi nếu có
    if ws_out.max_row >= 3:
        ws_out.delete_rows(3, ws_out.max_row - 2)

    today_dt = datetime.date.today()
    today_str = today_dt.strftime("%d/%m/%Y")
    out_row = 3
    po_codes = set()
    total_qty = 0
    skipped_zero_count = 0

    for r in range(2, ws_po.max_row + 1):
        po_code = ws_po.cell(r, col_map['po']).value
        product_code = ws_po.cell(r, col_map['product']).value
        qty_val = ws_po.cell(r, col_map['qty']).value
        est_date_val = ws_po.cell(r, col_map['est_date']).value
        supplier_val = ws_po.cell(r, col_map['supplier']).value
        nsx_val = ws_po.cell(r, col_map['nsx']).value
        hsd_val = ws_po.cell(r, col_map['hsd']).value

        # Bỏ qua nếu cả mã PO và mã hàng đều trống
        if po_code is None and product_code is None:
            continue

        # Kiểm tra số lượng thực nhận ở cột AD
        try:
            qty_num = float(qty_val) if qty_val is not None else 0
        except (ValueError, TypeError):
            qty_num = 0

        # LƯU Ý: Nếu số lượng thực nhận = 0 (hoặc <= 0) thì bỏ luôn dòng đó
        if qty_num <= 0:
            skipped_zero_count += 1
            continue

        if po_code:
            po_codes.add(str(po_code).strip())

        final_qty = int(qty_num) if qty_num.is_integer() else qty_num
        total_qty += final_qty

        # Cột A: orderInboundCode <- Mã PO (Cột B)
        ws_out.cell(out_row, 1, str(po_code).strip() if po_code is not None else "")
        # Cột B: productCode <- Mã hàng (Cột T)
        ws_out.cell(out_row, 2, str(product_code).strip() if product_code is not None else "")
        # Cột C: expectedQuantity <- Số lượng PR thực nhận (Cột AD)
        ws_out.cell(out_row, 3, final_qty)
        # Cột D: estimateReceiveTime <- Ngày giao hàng NCC xác nhận (Cột G), kiểu text '@', nếu < today thì đổi thành today
        cell_d = ws_out.cell(out_row, 4, process_estimate_date(est_date_val, today_dt))
        cell_d.number_format = '@'
        # Cột E: customerNote <- Để trống (hoặc ghi chú tùy chọn nếu có)
        ws_out.cell(out_row, 5, note if note else "")
        # Cột F: zoneType <- Điền B2B
        ws_out.cell(out_row, 6, zone)
        # Cột G: supplier <- Tên nhà cung cấp (Cột K)
        ws_out.cell(out_row, 7, str(supplier_val).strip() if supplier_val is not None else "")
        # Cột H: productionDate <- NSX (Cột X)
        cell_h = ws_out.cell(out_row, 8, format_date_str(nsx_val))
        cell_h.number_format = '@'
        # Cột I: expiryDate <- HSD (Cột Y)
        cell_i = ws_out.cell(out_row, 9, format_date_str(hsd_val))
        cell_i.number_format = '@'
        # Cột J: inboundDate <- Today (Ngày hôm nay)
        cell_j = ws_out.cell(out_row, 10, today_str)
        cell_j.number_format = '@'

        out_row += 1

    # Đảm bảo định dạng cột D, H, I, J là text (@)
    ws_out.column_dimensions['D'].number_format = '@'
    ws_out.column_dimensions['H'].number_format = '@'
    ws_out.column_dimensions['I'].number_format = '@'
    ws_out.column_dimensions['J'].number_format = '@'

    total_converted = out_row - 3
    if total_converted == 0:
        print("Cảnh báo: Không có dòng dữ liệu hợp lệ nào được chuyển đổi.")
        return False

    if not output_path:
        first_po = list(po_codes)[0] if po_codes else "DATA"
        date_stamp = datetime.datetime.now().strftime("%Y%m%d")
        output_path = f"file_nhap_hang_{first_po}_{date_stamp}.xlsx"

    wb_out.save(output_path)
    print("=" * 60)
    print("CHUYỂN ĐỔI THÀNH CÔNG!")
    print(f"- Số dòng nhập hàng: {total_converted}")
    print(f"- Đã loại bỏ số dòng có SL thực nhận = 0: {skipped_zero_count}")
    print(f"- Mã PO: {', '.join(sorted(po_codes))}")
    print(f"- Tổng số lượng nhập: {total_qty}")
    print(f"- File đã xuất: {output_path}")
    print("=" * 60)
    return True

if __name__ == "__main__":
    po_arg = sys.argv[1] if len(sys.argv) > 1 else None
    out_arg = sys.argv[2] if len(sys.argv) > 2 else None
    template_arg = sys.argv[3] if len(sys.argv) > 3 else None
    convert_po_to_inbound(po_arg, template_path=template_arg, output_path=out_arg)
